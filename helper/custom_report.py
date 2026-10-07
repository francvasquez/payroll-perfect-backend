"""
Custom Report: stack selected TIME CARDS / PAYROLL tables from stored
results.json across pay periods into one XLSX (one sheet per table).
"""

from __future__ import annotations

import io
import json
import logging
from datetime import datetime, timezone

import boto3
import openpyxl
from botocore.exceptions import ClientError
from openpyxl.styles import Font

import app_config
from app_config import S3_BUCKET
from exceptions import AppError

logger = logging.getLogger()
logger.setLevel(logging.INFO)

s3_client = boto3.client("s3")

# Matches filter_and_sort_df_to_dict(max_rows=200) in helper/results.py
RESULTS_ROW_CAP = 200

PAY_DATE_HEADER = "Pay Date"

# UI titles + fallback headers (post-rename, as stored in results.json).
# Keep in sync with pp-react-app TA_TABLE_OPTIONS / WFN_TABLE_OPTIONS and results.py.
def _renamed(cols, rename_map=None):
    rename_map = rename_map or {}
    return [rename_map.get(c, c) for c in cols]


TABLE_CATALOG = {
    ("ta", "break_credit_summary"): {
        "title": "Break Credit Summary",
        "headers": _renamed(app_config.COLS_ANOMALIES),
    },
    ("ta", "short_break_earned_credits"): {
        "title": "Short Break",
        "headers": _renamed(
            app_config.COLS_PRINT3a,
            {"Regular Rate Paid": "Straight Rate ($)"},
        ),
    },
    ("ta", "did_not_break_meal_waiver_check"): {
        "title": "Meal Waiver Check",
        "headers": _renamed(app_config.COLS_PRINT2_B),
    },
    ("ta", "ot_vs_wfn"): {
        "title": "Overtime on Time Card vs Payroll Check",
        "headers": _renamed(app_config.COLS_PRINT9),
    },
    ("ta", "dt_vs_wfn"): {
        "title": "Doubletime on Time Card vs Payroll Check",
        "headers": _renamed(app_config.COLS_PRINT9a),
    },
    ("ta", "seven_consecutive"): {
        "title": "Consecutive Days Check",
        "headers": _renamed(
            app_config.COLS_PRINT8,
            {"Attributed_Workday": "Trigger Date"},
        ),
    },
    ("ta", "split_shift"): {
        "title": "Split Shift Check",
        "headers": _renamed(
            app_config.COLS_PRINT5,
            {"Regular Rate Paid": "Straight Rate ($)"},
        ),
    },
    ("ta", "short_shift"): {
        "title": "Reporting Time Pay Warning",
        "headers": _renamed(app_config.COLS_PRINT3b),
    },
    ("wfn", "overtime_checks_variances"): {
        "title": "Overtime RROP vs Actual Paid",
        "headers": _renamed(
            app_config.COLUMNS_TO_SHOW,
            {
                "Variance": "Variance ($)",
                "1.5 OT Earnings Due": "1.5 OT Earnings Due ($)",
                "Actual Pay Check": "Actual Pay Check ($)",
            },
        ),
    },
    ("wfn", "doubletime_checks_variances"): {
        "title": "Doubletime RROP vs Actual Paid",
        "headers": _renamed(
            app_config.COLUMNS_TO_SHOW_DBLE,
            {
                "Double Time Due": "Double Time Due ($)",
                "Actual Pay Check Double": "Actual Pay Check Double ($)",
                "Actual Pay Check Dble": "Actual Pay Check Dble ($)",
                "Variance Dble": "Variance Dble ($)",
            },
        ),
    },
    ("wfn", "break_credit_variances"): {
        "title": "Break Credit RROP vs Actual Paid",
        "headers": _renamed(
            app_config.COLUMNS_TO_SHOW_BRKCRD,
            {
                "Actual Pay BrkCrd": "Actual Paid Break Credit",
                "Variance BrkCrd": "Variance Break Credit",
            },
        ),
    },
    ("wfn", "rest_credit_variances"): {
        "title": "Rest Credit RROP vs Actual Paid",
        "headers": _renamed(
            app_config.COLUMNS_TO_SHOW_REST,
            {
                "Actual Pay RestCrd": "Actual Paid Rest Credit ($)",
                "Variance RestCrd": "Variance Rest Credit ($)",
            },
        ),
    },
    ("wfn", "sick_credit_variances"): {
        "title": "Sick RROP vs Actual Paid",
        "headers": _renamed(
            app_config.COLUMNS_TO_SHOW_SICK,
            {
                "Sick Credit Due": "Sick Credit Due ($)",
                "Sick Paid": "Actual Paid Sick Credit ($)",
                "Variance Sick": "Variance Sick Credit ($)",
                "Regular Rate Paid": "Regular Rate Paid ($)",
            },
        ),
    },
    ("wfn", "flsa_check"): {
        "title": "FLSA Check",
        "headers": _renamed(app_config.COLUMNS_TO_SHOW_FLSA),
    },
    ("wfn", "min_wage_check"): {
        "title": "Minimum Wage Check",
        "headers": _renamed(app_config.COLUMNS_TO_SHOW_MINWAGE),
    },
    ("wfn", "non_active_check"): {
        "title": "Non-Active Check",
        "headers": _renamed(
            app_config.COLUMNS_TO_SHOW_NONACTIVE,
            {
                "Regular Hours": "Straight Hours Worked",
                "Regular Rate Paid": "Regular Rate Paid ($)",
            },
        ),
    },
}


def _blank_if_na(value):
    if value is None:
        return ""
    if isinstance(value, float) and value != value:
        return ""
    return value


def _sheet_title(title: str, used: set[str]) -> str:
    """Excel sheet names max 31 chars; cannot contain \\ / ? * [ ]."""
    base = str(title or "Sheet").strip() or "Sheet"
    for ch in ("\\", "/", "?", "*", "[", "]"):
        base = base.replace(ch, "-")
    base = base[:31] or "Sheet"
    sheet = base
    n = 2
    while sheet in used:
        suffix = f"_{n}"
        sheet = f"{base[: 31 - len(suffix)]}{suffix}"
        n += 1
    used.add(sheet)
    return sheet


def _normalize_tables(tables) -> list[dict]:
    if not tables:
        raise AppError("Select at least one table.", status_code=400)

    normalized = []
    seen = set()
    for item in tables:
        if not isinstance(item, dict):
            raise AppError("Each table selection must be an object.", status_code=400)
        section = str(item.get("section") or "").strip().lower()
        key = str(item.get("key") or "").strip()
        title = str(item.get("title") or "").strip()
        if section not in ("ta", "wfn") or not key:
            raise AppError(
                "Each table needs section ('ta' or 'wfn') and key.",
                status_code=400,
            )
        catalog = TABLE_CATALOG.get((section, key))
        if not catalog:
            raise AppError(f"Unknown table: {section}.{key}", status_code=400)
        if not title:
            title = catalog["title"]
        dedupe = (section, key)
        if dedupe in seen:
            continue
        seen.add(dedupe)
        normalized.append(
            {
                "section": section,
                "key": key,
                "title": title,
                "fallback_headers": list(catalog["headers"]),
            }
        )
    if not normalized:
        raise AppError("Select at least one table.", status_code=400)
    return normalized


def _load_results_json(client_id: str, pay_date: str) -> dict | None:
    key = f"clients/{client_id}/processed/{pay_date}/results.json"
    try:
        response = s3_client.get_object(Bucket=S3_BUCKET, Key=key)
        return json.loads(response["Body"].read())
    except ClientError as e:
        error_code = e.response.get("Error", {}).get("Code")
        if error_code in ("NoSuchKey", "404"):
            return None
        logger.exception("Failed loading results.json for %s/%s", client_id, pay_date)
        raise AppError(
            f"Failed to load results for {pay_date}.",
            status_code=500,
        )


def _collect_headers(rows: list[dict], fallback: list[str]) -> list[str]:
    """Pay Date first, then fallback order, then any extra keys seen in rows."""
    ordered: list[str] = [PAY_DATE_HEADER]
    seen = {PAY_DATE_HEADER}
    for col in fallback:
        if col not in seen:
            ordered.append(col)
            seen.add(col)
    for row in rows:
        for col in row.keys():
            if col not in seen:
                ordered.append(col)
                seen.add(col)
    return ordered


def _write_sheet(ws, headers: list[str], rows: list[dict]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([_blank_if_na(row.get(h)) for h in headers])


def _build_workbook(table_payloads: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    used_titles: set[str] = set()
    for payload in table_payloads:
        ws = wb.create_sheet(_sheet_title(payload["title"], used_titles))
        headers = _collect_headers(payload["rows"], payload["fallback_headers"])
        _write_sheet(ws, headers, payload["rows"])

    if not wb.worksheets:
        ws = wb.create_sheet("Custom Report")
        _write_sheet(ws, [PAY_DATE_HEADER], [])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_custom_report(client_id: str, pay_dates: list, tables) -> dict:
    if not client_id:
        raise AppError("clientId is required.", status_code=400)
    if not pay_dates:
        raise AppError("Select at least one pay period.", status_code=400)

    ordered_dates = sorted({str(d) for d in pay_dates})
    selected_tables = _normalize_tables(tables)

    # Accumulate rows per selected table (preserve selection order)
    stacked: dict[tuple[str, str], list[dict]] = {
        (t["section"], t["key"]): [] for t in selected_tables
    }
    warnings: list[str] = []
    capped: list[dict] = []
    periods_loaded = 0

    for pay_date in ordered_dates:
        results = _load_results_json(client_id, pay_date)
        if results is None:
            warnings.append(f"{pay_date}: not a processed pay period (missing results.json).")
            continue

        periods_loaded += 1
        section_blobs = {
            "ta": results.get("ta") or {},
            "wfn": results.get("wfn") or {},
        }

        for table in selected_tables:
            section = table["section"]
            key = table["key"]
            raw = section_blobs.get(section, {}).get(key)
            rows = raw if isinstance(raw, list) else []

            if len(rows) >= RESULTS_ROW_CAP:
                capped.append(
                    {
                        "payDate": pay_date,
                        "section": section,
                        "key": key,
                        "title": table["title"],
                        "rowCount": len(rows),
                    }
                )
                warnings.append(
                    f"{pay_date} / {table['title']}: stored results hit the "
                    f"{RESULTS_ROW_CAP}-row display cap — this sheet may be incomplete."
                )

            for row in rows:
                if not isinstance(row, dict):
                    continue
                out = {PAY_DATE_HEADER: pay_date}
                out.update(row)
                stacked[(section, key)].append(out)

    if periods_loaded == 0:
        raise AppError(
            "Custom report could not load any periods. " + " | ".join(warnings),
            status_code=400,
        )

    table_payloads = [
        {
            "title": t["title"],
            "fallback_headers": t["fallback_headers"],
            "rows": stacked[(t["section"], t["key"])],
        }
        for t in selected_tables
    ]

    xlsx_bytes = _build_workbook(table_payloads)

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    filename = f"custom_report_{client_id}_{stamp}.xlsx"
    s3_key = f"clients/{client_id}/exports/{filename}"

    s3_client.put_object(
        Bucket=S3_BUCKET,
        Key=s3_key,
        Body=xlsx_bytes,
        ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    download_url = s3_client.generate_presigned_url(
        "get_object",
        Params={"Bucket": S3_BUCKET, "Key": s3_key},
        ExpiresIn=600,
    )

    total_rows = sum(len(p["rows"]) for p in table_payloads)

    return {
        "downloadUrl": download_url,
        "s3Key": s3_key,
        "filename": filename,
        "rowCount": total_rows,
        "periodCount": periods_loaded,
        "tableCount": len(selected_tables),
        "rowCap": RESULTS_ROW_CAP,
        "cappedTables": capped,
        "incomplete": len(capped) > 0,
        "warnings": warnings,
    }
